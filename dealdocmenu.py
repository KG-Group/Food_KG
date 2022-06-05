from docx import Document
import xlrd
import xlwt
document = Document('./document/OlymMenu.docx')
def fil(s):
	return s if s != '' else None
DishList = []
Type = ''
def samechr(s1,s2):
	cnt = 0
	for chr1 in s1:
		for chr2 in s2:
			if(chr1 == chr2):
				cnt +=1
				break
	return cnt
for p in document.paragraphs:
	s = ['','']
	s[0] = p.text.split('\t')[0]
	s[1] = p.text.split('\t')[-1]
	if(s[0] == s[1]):
		Type = s[0]
	else:
		ThisDish = [s[0],Type,s[1]]
		DishList.append(ThisDish)


QuarterMenuList = {}
data = xlrd.open_workbook("./document/QuarterRowMenu.xls")
table = data.sheets()[0]

for i in range(0,24154):
    nowRow = table.row_values(i,start_colx = 0,end_colx = 50)
    nowRow = [item for item in nowRow if item != '']
    QuarterMenuList[nowRow[0]] = nowRow[1:]
cnt = 0
for i in range (0,len(DishList)):
	good = 0
	best = 0
	print(i)
	for Menu in QuarterMenuList.keys():
		if(str(DishList[i][0]) == str(Menu)):
			DishList[i].append(Menu)
			DishList[i].append(QuarterMenuList[Menu])
			best = 1
			break
	if (best == 0):
		for Menu in QuarterMenuList.keys():
			if(str(DishList[i][0]) in str(Menu)):
				DishList[i].append(Menu)
				DishList[i].append(QuarterMenuList[Menu])
				good = 1
				break
	if(good == 0):
		bestkey = list(QuarterMenuList.keys())[0]
		maxsame = samechr(str(DishList[0][0]),bestkey)
		for Menu in QuarterMenuList.keys():
			if(samechr(str(DishList[i][0]),str(Menu))>maxsame):
				bestkey = Menu
				maxsame = samechr(str(DishList[i][0]),str(Menu))
		DishList[i].append(bestkey)
		DishList[i].append(QuarterMenuList[bestkey])
	print(DishList[i])

xl = xlwt.Workbook(encoding = 'utf-8')
sheet = xl.add_sheet("sheet1")
i = 0
for i in range(0,len(DishList)):
    j = 4
    sheet.write(i,0,DishList[i][0])
    sheet.write(i,1,DishList[i][1])
    sheet.write(i,2,DishList[i][2])
    sheet.write(i,3,DishList[i][3])
    for Mat in DishList[i][4]:
        sheet.write(i,j,Mat)
        j+=1
    i+=1
xl.save('./document/Oly2Menu.xls')